import docx
import json
import sys
import re

sys.stdout.reconfigure(encoding='utf-8')

# 1. Extrair logistica completa
doc_log = docx.Document(r'projects/analise-logistica-ze041-vitoria-do-mearim/documentos da zona 41/logistica-completa-vitoria-do-mearim-2026.docx')

# Tabela 2: 7 Rotas de Veiculos
rotas_veiculos = []
t_veic = doc_log.tables[1]
headers_v = [c.text.strip() for c in t_veic.rows[0].cells]
for r in t_veic.rows[1:]:
    vals = [c.text.strip().replace('\n', ' ') for c in r.cells]
    if len(vals) >= 9:
        rotas_veiculos.append({
            "rota": int(vals[0]) if vals[0].isdigit() else vals[0],
            "veiculo_modelo_placa": vals[1],
            "motorista_telefone": vals[2],
            "itinerario": vals[3],
            "qtd_locais": int(vals[4]) if vals[4].isdigit() else vals[4],
            "qtd_urnas_cabinas": int(vals[5]) if vals[5].isdigit() else vals[5],
            "area": vals[6],
            "distancia_total": vals[7],
            "tempo_estimado": vals[8]
        })

# Tabela 4: 31 Rotas de Moto
rotas_motos = []
t_motos = doc_log.tables[3]
for r in t_motos.rows[1:]:
    vals = [c.text.strip().replace('\n', ' ') for c in r.cells]
    if len(vals) >= 8:
        rotas_motos.append({
            "moto_num": int(vals[0]) if vals[0].isdigit() else vals[0],
            "condutor_telefone": vals[1],
            "veiculo_placa_ano": vals[2],
            "local_atendido": vals[3],
            "qtd_urnas": int(vals[4]) if vals[4].isdigit() else vals[4],
            "area": vals[5],
            "distancia_total": vals[6],
            "tempo_estimado": vals[7]
        })

# 2. Extrair Administradores de Prédio
doc_adm = docx.Document(r'projects/analise-logistica-ze041-vitoria-do-mearim/documentos da zona 41/CONTROLE DE PRESENÇA ADM DE PREDIO.docx')
adms = []
for t in doc_adm.tables:
    for r in t.rows:
        vals = [c.text.strip().replace('\n', ' ') for c in r.cells]
        if len(vals) >= 3 and vals[0].isdigit():
            adms.append({
                "num": int(vals[0]),
                "nome": vals[1],
                "local": vals[2]
            })

# 3. Extrair Coordenadores de Acessibilidade
doc_acess = docx.Document(r'projects/analise-logistica-ze041-vitoria-do-mearim/documentos da zona 41/CONTROLE DE PRESENÇA DE COORDENADORES DE ACESSIBILIDADE.docx')
coord_acess = []
for t in doc_acess.tables:
    for r in t.rows:
        vals = [c.text.strip().replace('\n', ' ') for c in r.cells]
        if len(vals) >= 3 and vals[0].isdigit():
            coord_acess.append({
                "num": int(vals[0]),
                "nome": vals[1],
                "local": vals[2]
            })

consolidado_log = {
    "edital_vigente": "Edital nº 20/2026 (Pós-2ª Vistoria)",
    "rotas_veiculos_carga": rotas_veiculos,
    "rotas_mototaxis": rotas_motos,
    "administradores_predio": adms,
    "coordenadores_acessibilidade": coord_acess
}

with open(r'projects/analise-logistica-ze041-vitoria-do-mearim/data/logistica_consolidada_2026.json', 'w', encoding='utf-8') as f:
    json.dump(consolidado_log, f, ensure_ascii=False, indent=2)

print(f"Sucesso!")
print(f"Rotas de Veículos (Carga): {len(rotas_veiculos)}")
print(f"Rotas de Mototáxis: {len(rotas_motos)}")
print(f"Administradores de Prédio: {len(adms)}")
print(f"Coordenadores de Acessibilidade: {len(coord_acess)}")
