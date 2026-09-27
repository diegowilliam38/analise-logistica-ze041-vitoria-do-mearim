import sys
import os
import json
import docx
import zipfile
import xml.etree.ElementTree as ET

sys.stdout.reconfigure(encoding='utf-8')

project_dir = r"C:\Users\Diego William\OneDrive\Área de Trabalho\AIOS-WORKSPACE\projects\analise-logistica-ze041-vitoria-do-mearim"
docs_dir = os.path.join(project_dir, "documentos da zona 41")

# 1. Checar CONTROLE DE PRESENÇA ADM DE PREDIO.docx
adm_file = os.path.join(docs_dir, "CONTROLE DE PRESENÇA ADM DE PREDIO.docx")
doc_adm = docx.Document(adm_file)
print("=== CONTROLE DE PRESENCA ADM DE PREDIO.docx ===")
print("Paragrafos:", len(doc_adm.paragraphs))
for p in doc_adm.paragraphs:
    if p.text.strip():
        print("  P:", p.text.strip())

print("Tabelas:", len(doc_adm.tables))
adms_extraidos = []
for t_idx, table in enumerate(doc_adm.tables):
    print(f"Tabela {t_idx} ({len(table.rows)} linhas):")
    for r_idx, row in enumerate(table.rows):
        cells = [c.text.strip().replace('\n', ' ') for c in row.cells]
        print(f"  L{r_idx}: {cells}")

# 2. Checar CONTROLE DE PRESENÇA DE COORDENADORES DE ACESSIBILIDADE.docx
coord_file = os.path.join(docs_dir, "CONTROLE DE PRESENÇA DE COORDENADORES DE ACESSIBILIDADE.docx")
doc_coord = docx.Document(coord_file)
print("\n=== CONTROLE DE PRESENCA DE COORDENADORES DE ACESSIBILIDADE.docx ===")
print("Paragrafos:", len(doc_coord.paragraphs))
for p in doc_coord.paragraphs:
    if p.text.strip():
        print("  P:", p.text.strip())

for t_idx, table in enumerate(doc_coord.tables):
    print(f"Tabela {t_idx} ({len(table.rows)} linhas):")
    for r_idx, row in enumerate(table.rows):
        cells = [c.text.strip().replace('\n', ' ') for c in row.cells]
        print(f"  L{r_idx}: {cells}")

# 3. Investigar CONTROLE_PRESENCA_MESARIOS_DIA_DA_ELEICAO.odt por seções
odt_path = os.path.join(docs_dir, "CONTROLE_PRESENCA_MESARIOS_DIA_DA_ELEICAO.odt")
with zipfile.ZipFile(odt_path) as z:
    xml_content = z.read('content.xml')

root = ET.fromstring(xml_content)
ns = {'table': 'urn:oasis:names:tc:opendocument:xmlns:table:1.0', 'text': 'urn:oasis:names:tc:opendocument:xmlns:text:1.0'}

tables = root.findall('.//table:table', ns)
print(f"\n=== CONTROLE_PRESENCA_MESARIOS_DIA_DA_ELEICAO.odt ===")
print(f"Total de tabelas ODT: {len(tables)}")

for t_i, table in enumerate(tables):
    rows = table.findall('.//table:table-row', ns)
    print(f"Tabela {t_i}: {len(rows)} linhas")
